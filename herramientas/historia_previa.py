"""Importación local y prudente de datos desde historias clínicas previas."""

from datetime import date, datetime
from io import BytesIO
import re
import unicodedata

import streamlit as st
from docx import Document
from pypdf import PdfReader


TIPOS_DOCUMENTO = {"NV", "RC", "TI", "CC", "CE", "PEP", "PS", "OTRO"}


def _normalizar(texto):
    texto = unicodedata.normalize("NFD", str(texto or ""))
    texto = "".join(caracter for caracter in texto if unicodedata.category(caracter) != "Mn")
    return re.sub(r"\s+", " ", texto).strip().upper()


def _texto_docx(datos):
    documento = Document(BytesIO(datos))
    partes = [parrafo.text.strip() for parrafo in documento.paragraphs if parrafo.text.strip()]
    for tabla in documento.tables:
        for fila in tabla.rows:
            valores = [celda.text.strip() for celda in fila.cells if celda.text.strip()]
            if valores:
                partes.append(" | ".join(valores))
    return "\n".join(partes)


@st.cache_data(show_spinner=False, max_entries=12)
def extraer_texto_historia_previa(nombre_archivo, datos):
    """Extrae texto de PDF con capa textual o de archivos DOCX.

    No guarda el documento ni lo envía a servicios externos. Los PDF escaneados
    sin capa de texto se reportan para que el usuario pueda transcribir los datos.
    """
    extension = nombre_archivo.rsplit(".", 1)[-1].lower() if "." in nombre_archivo else ""
    try:
        if extension == "pdf":
            lector = PdfReader(BytesIO(datos))
            texto = "\n".join((pagina.extract_text() or "") for pagina in lector.pages).strip()
            return texto or "__PDF_SIN_TEXTO_EXTRAIBLE__"
        if extension == "docx":
            return _texto_docx(datos)
    except Exception:
        return "__ARCHIVO_NO_LEIBLE__"
    return ""


def _valor_etiqueta(texto, etiquetas):
    for etiqueta in etiquetas:
        patron = rf"(?:^|\n)\s*{etiqueta}\s*[:\-]?\s*([^\n|]+)"
        coincidencia = re.search(patron, texto, flags=re.IGNORECASE)
        if coincidencia:
            valor = coincidencia.group(1).strip(" -:\t")
            if valor and _normalizar(valor) not in {_normalizar(etiqueta), "NO REGISTRADO"}:
                return valor
    return ""


def _fecha_desde_texto(valor):
    coincidencia = re.search(r"\b(\d{1,2})[/-](\d{1,2})[/-](\d{4})\b", str(valor or ""))
    if not coincidencia:
        return None
    try:
        return datetime.strptime("/".join(coincidencia.groups()), "%d/%m/%Y").date()
    except ValueError:
        return None


def _extraer_antecedentes(texto):
    lineas = [linea.strip() for linea in str(texto or "").splitlines()]
    inicio = None
    for indice, linea in enumerate(lineas):
        normalizada = _normalizar(linea).rstrip(":")
        if normalizada in {
            "ANTECEDENTES",
            "ANTECEDENTES GENERALES",
            "ANTECEDENTES PERSONALES Y FAMILIARES",
        }:
            inicio = indice + 1
            break
    if inicio is None:
        return ""

    cierres = {
        "NEURODESARROLLO", "REVISION POR SISTEMAS", "SIGNOS VITALES",
        "EXAMEN FISICO", "PARACLINICOS", "ANALISIS", "DIAGNOSTICOS",
        "IMPRESION DIAGNOSTICA", "PLAN", "CONDUCTA FINAL", "EVOLUCION",
    }
    bloque = []
    for linea in lineas[inicio:]:
        normalizada = _normalizar(linea).rstrip(":")
        if normalizada in cierres:
            break
        if linea:
            bloque.append(linea)
    return "\n".join(bloque).strip()


def _limpiar_informante(valor):
    valor = str(valor or "").strip()
    return re.split(r"\s*/\s*(?:TEL[EÉ]FONO|CELULAR|PROCEDENTE|EPS)\s*[:\-]", valor, maxsplit=1, flags=re.IGNORECASE)[0].strip()


def _formatear_informante(valor):
    """Conserva solo parentesco, edad y ocupación, sin inventar parentescos."""
    valor = _limpiar_informante(valor)
    if not valor:
        return ""

    parentesco = re.search(
        r"\b(MADRE|PADRE|ABUELA|ABUELO|HERMANA|HERMANO|TIA|TIO|CUIDADOR(?:A)?|ACUDIENTE)\b",
        _normalizar(valor),
    )
    edad = re.search(
        r"\b(?:EDAD\s*[:\-]?\s*)?(\d{1,3})\s*A[NÑ]OS?\b",
        valor,
        flags=re.IGNORECASE,
    )
    ocupacion = re.search(r"\bOCUPACI[ÓO]N\s*[:\-]?\s*([^/|\n]+)", valor, flags=re.IGNORECASE)
    if not ocupacion:
        ocupacion = re.search(r"\b\d{1,3}\s*A[NÑ]OS?\s*[-/]\s*([^/|\n]+)", valor, flags=re.IGNORECASE)

    partes = [parentesco.group(1) if parentesco else "INFORMANTE"]
    if edad:
        partes.append(f"{edad.group(1)} AÑOS")
    if ocupacion and ocupacion.group(1).strip():
        partes.append(ocupacion.group(1).strip(" .-"))

    # Si no hay ningún dato estructurado, conservar el texto original para revisión manual.
    return " - ".join(partes) if len(partes) > 1 or parentesco else valor


def _tipo_documento_pediatrico(fecha_nacimiento):
    """Aplica RC/TI solo a menores cuando el documento previo no lo informa."""
    if not fecha_nacimiento:
        return ""
    hoy = date.today()
    edad = hoy.year - fecha_nacimiento.year - (
        (hoy.month, hoy.day) < (fecha_nacimiento.month, fecha_nacimiento.day)
    )
    if not 0 <= edad < 18:
        return ""
    return "RC" if edad < 7 else "TI"


def _normalizar_tipo_documento(valor):
    valor = _normalizar(valor)
    equivalencias = {
        "REGISTRO CIVIL": "RC",
        "TARJETA DE IDENTIDAD": "TI",
        "TARJETA IDENTIDAD": "TI",
        "CEDULA DE CIUDADANIA": "CC",
        "CEDULA EXTRANJERIA": "CE",
    }
    return equivalencias.get(valor, valor if valor in TIPOS_DOCUMENTO else "")


def _valor_en_linea(texto, etiqueta):
    coincidencia = re.search(rf"\b(?:{etiqueta})\s*[:\-]\s*([^\n/|]+)", str(texto or ""), flags=re.IGNORECASE)
    return coincidencia.group(1).strip() if coincidencia else ""


def extraer_datos_historia_previa(texto):
    """Obtiene únicamente datos de continuidad que pueden migrarse con seguridad."""
    texto = str(texto or "")[:120000]
    fecha_nacimiento = _fecha_desde_texto(_valor_etiqueta(
        texto, [r"FECHA DE NACIMIENTO", r"FECHA NACIMIENTO", r"FN"]
    ))
    tipo_documento = _normalizar_tipo_documento(_valor_etiqueta(
        texto, [r"TIPO DE DOCUMENTO", r"TIPO DOCUMENTO"]
    ))
    if not tipo_documento:
        tipo_documento = _tipo_documento_pediatrico(fecha_nacimiento)
    sexo = _valor_etiqueta(texto, [r"SEXO"])
    sexo_normalizado = _normalizar(sexo)
    if sexo_normalizado.startswith("FEMEN"):
        sexo = "Femenino"
    elif sexo_normalizado.startswith("MASCUL"):
        sexo = "Masculino"
    else:
        sexo = ""

    return {
        "nombre": _valor_etiqueta(texto, [r"NOMBRES? Y APELLIDOS?", r"NOMBRE(?: DEL)?(?: PACIENTE| RN)?"]),
        "tipo_documento": tipo_documento,
        "documento": _valor_etiqueta(texto, [r"DOCUMENTO", r"IDENTIFICACION", r"N[ÚU]MERO DE DOCUMENTO"]),
        "fecha_nacimiento": fecha_nacimiento,
        "sexo": sexo,
        "eps": _valor_etiqueta(texto, [r"EPS", r"ASEGURADORA"]),
        "telefono": _valor_etiqueta(texto, [r"TEL[EÉ]FONO", r"CELULAR"])
        or _valor_en_linea(texto, r"TEL[EÉ]FONO|CELULAR"),
        "informante": _formatear_informante(_valor_etiqueta(
            texto,
            [r"INFORMANTE(?:\s*\([^\n)]*\))?(?:\s*/\s*ACOMPA[ÑN]ANTE)?", r"ACOMPA[ÑN]ANTE"],
        )),
        "antecedentes": _extraer_antecedentes(texto),
    }


def _resumen_datos(datos):
    etiquetas = {
        "nombre": "Nombre", "tipo_documento": "Tipo de documento", "documento": "Documento",
        "fecha_nacimiento": "Fecha de nacimiento", "sexo": "Sexo", "eps": "EPS",
        "telefono": "Teléfono", "informante": "Informante", "antecedentes": "Antecedentes",
    }
    partes = []
    for clave, etiqueta in etiquetas.items():
        valor = datos.get(clave)
        if clave == "fecha_nacimiento" and valor:
            valor = valor.strftime("%d/%m/%Y")
        if valor:
            if clave == "antecedentes":
                partes.append(f"{etiqueta}: se identificó un bloque clínico editable.")
            else:
                partes.append(f"{etiqueta}: {valor}")
    return "\n".join(partes) or "No se identificaron datos estructurados para migrar."


def render_importador_historia_previa(
    prefix,
    campos,
    antecedentes_default="",
    campos_protegidos=(),
    campos_reset_por_migracion=(),
):
    """Renderiza un importador reutilizable antes de identificación.

    `campos` relaciona nombres estándar con las claves de Session State de cada
    plantilla. Solo se migra información identificada de forma explícita.
    """
    with st.expander("Historia clínica previa", expanded=False):
        st.caption(
            "Adjunte una historia previa en PDF con texto o Word (.docx). Se lee localmente "
            "para migrar identificación y antecedentes; el documento no se adjunta al informe actual."
        )
        archivo = st.file_uploader(
            "Seleccionar historia previa",
            type=["pdf", "docx"],
            accept_multiple_files=False,
            key=f"{prefix}_historia_previa_archivo",
        )
        if not archivo:
            return

        texto = extraer_texto_historia_previa(archivo.name, archivo.getvalue())
        if texto == "__PDF_SIN_TEXTO_EXTRAIBLE__":
            st.warning("El PDF parece escaneado y no contiene texto seleccionable. Use un PDF con texto o un archivo Word (.docx).")
            return
        if texto in {"", "__ARCHIVO_NO_LEIBLE__"}:
            st.error("No fue posible leer el archivo. Verifique que sea un PDF o Word (.docx) válido.")
            return

        datos = extraer_datos_historia_previa(texto)

        st.text_area(
            "Datos preparados para migración",
            value=_resumen_datos(datos),
            height=170,
            disabled=True,
        )
        col_migracion_1, col_migracion_2 = st.columns(2)
        actualizar_identificacion = col_migracion_1.checkbox(
            "Actualizar identificación del paciente",
            value=True,
            key=f"{prefix}_historia_previa_actualizar_identificacion",
        )
        actualizar_antecedentes = col_migracion_2.checkbox(
            "Actualizar antecedentes documentados",
            value=True,
            key=f"{prefix}_historia_previa_actualizar_antecedentes",
        )
        st.caption(
            "No se modifican motivo de consulta, enfermedad actual, revisión, signos vitales, examen físico, "
            "diagnósticos, análisis ni plan del ejemplo seleccionado."
        )
        if st.button(
            "Actualizar datos desde historia previa",
            key=f"{prefix}_migrar_historia_previa",
            use_container_width=True,
        ):
            campos_reset = set(campos_reset_por_migracion)
            estado_protegido = {
                clave: st.session_state.get(clave)
                for clave in campos_protegidos
                if clave in st.session_state and clave not in campos_reset
            }
            migrados = []
            for campo, clave_estado in campos.items():
                valor = datos.get(campo)
                if not valor or not clave_estado:
                    continue
                if campo == "antecedentes" and not actualizar_antecedentes:
                    continue
                if campo != "antecedentes" and not actualizar_identificacion:
                    continue
                st.session_state[clave_estado] = valor
                migrados.append(campo)
            if migrados:
                # Los valores fisiológicos y antropométricos pertenecen a la atención actual,
                # nunca al ejemplo o al documento previo.
                for clave in campos_reset:
                    if clave in st.session_state:
                        st.session_state[clave] = ""
                # El rerun posterior no puede alterar el escenario clínico cargado.
                for clave, valor in estado_protegido.items():
                    st.session_state[clave] = valor
                st.session_state[f"{prefix}_historia_previa_notice"] = (
                    "Datos actualizados desde historia previa: " + ", ".join(migrados)
                    + ". El contenido clínico del ejemplo se conservó sin cambios. "
                    "Registre nuevamente los signos vitales y la antropometría medidos hoy."
                )
                st.rerun()
            st.info("No se modificó ningún campo: no se identificaron datos seleccionados para actualizar.")
